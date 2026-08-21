#!/usr/bin/env python3
"""
validate.py — structural check for a Be10X FAQ .docx.

House style = exactly 35 paragraphs:
  [0] eyebrow  [1] title  [2] subtitle  [3] accent rule  [4] source note
  then 15 × ( Q# heading , A answer )  = 30
All text runs must be Arial; the footer must carry PAGE/NUMPAGES fields.

    python validate.py <file.docx>      -> "All validations PASSED!" or a list of failures
"""
import sys
from pathlib import Path


def validate(path):
    from docx import Document
    doc = Document(str(path))
    paras = doc.paragraphs
    errors = []

    if len(paras) != 35:
        errors.append(f"expected 35 paragraphs, found {len(paras)}")

    def text(i):
        return paras[i].text.strip() if i < len(paras) else ""

    if "BE10X" not in text(0).upper():
        errors.append(f"[para 0] expected eyebrow 'BE10X…', got {text(0)[:40]!r}")
    if "Frequently Asked Questions" not in text(2):
        errors.append(f"[para 2] expected subtitle, got {text(2)[:40]!r}")

    qa = paras[5:]
    if len(qa) != 30:
        errors.append(f"expected 30 Q&A paragraphs (15 pairs), found {len(qa)}")
    else:
        for idx in range(15):
            q, a = qa[idx * 2], qa[idx * 2 + 1]
            if not q.text.strip().startswith(f"Q{idx + 1}."):
                errors.append(f"Q{idx + 1} heading malformed: {q.text[:50]!r}")
            if not a.text.strip().startswith("A."):
                errors.append(f"A{idx + 1} malformed: {a.text[:50]!r}")

    bad_fonts = {r.font.name for p in paras for r in p.runs if r.font.name and r.font.name != "Arial"}
    if bad_fonts:
        errors.append(f"non-Arial fonts present: {sorted(bad_fonts)}")

    footer_xml = doc.sections[0].footer.paragraphs[0]._p.xml
    if "PAGE" not in footer_xml or "NUMPAGES" not in footer_xml:
        errors.append("footer missing PAGE/NUMPAGES fields")

    return errors


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: python validate.py <file.docx>"); sys.exit(2)
    errs = validate(sys.argv[1])
    if errs:
        print("VALIDATION FAILED:")
        for e in errs:
            print(f"  - {e}")
        sys.exit(1)
    print("All validations PASSED!")
